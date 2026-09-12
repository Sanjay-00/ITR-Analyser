"""
excel_generator.py  -  the ITR Validation spreading sheet, in the combined
master format (see template_config's docstring).

Sheet 1, "ITR Validation", is the analyst's own layout and nothing else:
Profit & Loss, Balance Sheet, Ratios, DSCR Calculation and Financial Snap,
one column per financial year, amounts in lakhs. Only the figures read from
the accounts are written as values; every total, profit line, ratio and
Snap figure is a LIVE Excel formula (template_config.TOTAL_FORMULAS / RATIOS
/ SNAP), so the analyst can correct one input and everything recomputes.

A row the accounts do not use shows 0. A row that could not be READ (that
year's P&L or Balance Sheet did not reconcile) also shows 0 - so every
formula still works - but is filled red with a note, because "the business
has no income" and "we could not read its income" must never look alike.

Sheet 2, "Analysis", carries year-on-year trends and red/amber flags; sheet
3, "Audit Trail", where every figure came from and what could not be
verified.
"""

import io
import re
import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import analysis as A
from .extract import financials as F
from .mapping import taxonomy as T
from .mapping import template_config as C

# ─────────────────────────────────────────────
# PALETTE AND FORMATS  (matched to the analysts' reference sheets)
# ─────────────────────────────────────────────

NAVY       = "1F3864"
WHITE      = "FFFFFF"
HEAD_BG    = "DDEBF7"     # header rows and the RATIOS block
SECTION_BG = "D6E4F0"     # Analysis / Audit Trail section bands
TOTAL_BG   = "FFFF00"     # the analysts' yellow totals
UNREAD_BG  = "FFC7CE"     # a figure that could not be read from the ITR
BORDER_CLR = "808080"
GOOD_GREEN = "375623"
BAD_RED    = "C00000"
WARN_BG    = "FFF2F2"

AMT_FMT  = "0.00"
DATE_FMT = "[$-409]d\\-mmm\\-yy;@"
PCT_FMT  = "0.0%;-0.0%;0.0%"

LAKH = 100000.0

ROWS      = C.ROWS
SNAP_ROWS = [label for label, _f in C.SNAP]

UNREAD_NOTE = ("Not read from the ITR - this year's statement did not "
               "reconcile, so 0 is a placeholder. Please verify and enter "
               "the figure.")


def _b(style="thin"):
    s = Side(style=style, color=BORDER_CLR)
    return Border(left=s, right=s, top=s, bottom=s)


def _f(size=10, bold=False, color="000000", italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def _a(h="left", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def _formula(template: str, refs: dict) -> str:
    """'{gross_receipts}-{gross_expenses}' -> '=B7-B17' for one column."""
    return "=" + re.sub(r"\{(\w+)\}", lambda m: refs[m.group(1)], template)


# ─────────────────────────────────────────────
# ENTRY POINTS
# ─────────────────────────────────────────────

def borrower_name(columns: list) -> str:
    """
    The borrower's name for the title and filename.

    Taken from a column that actually produced statements. A year with none -
    a NIL return, or a document that failed to read - has no reliable name to
    offer, and letting it supply one titled a whole workbook after a fragment
    of ITR form boilerplate.
    """
    for cols in (
        [c for c in columns if c.get("blocks_used") and c.get("entity")],
        [c for c in columns if c.get("entity")],
    ):
        if cols:
            return cols[0]["entity"]
    return "Borrower"


def generate_excel(columns: list) -> bytes:
    """`columns` is columns.spread_many()'s output, oldest year first."""
    wb = Workbook()
    _build_sheet(wb, columns)
    _build_analysis_sheet(wb, columns)
    _build_audit_sheet(wb, columns)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _status(col: dict) -> str:
    return "Provisional" if col.get("provisional") else "Audited"


def _col_header(col: dict) -> str:
    y = col.get("year")
    return f"31-03-{y}" if y else (col.get("source_name") or "Unknown year")


# ─────────────────────────────────────────────
# SHEET 1 - ITR VALIDATION (the master format)
# ─────────────────────────────────────────────

def _period_header(ws, row: int, title, columns: list, merge_title: bool) -> None:
    """Two header rows: title | Audited/Provisional..., then (blank) | dates."""
    n = max(len(columns), 1)
    t = ws.cell(row=row, column=1, value=title)
    t.font, t.fill, t.border = _f(bold=True), _fill(HEAD_BG), _b()
    ws.cell(row=row + 1, column=1).border = _b()
    ws.cell(row=row + 1, column=1).fill = _fill(HEAD_BG)
    if merge_title:
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=1)
        t.alignment = _a(v="center", wrap=True)
    for i in range(n):
        col = columns[i] if i < len(columns) else {}
        h = ws.cell(row=row, column=2 + i, value=_status(col) if col else "")
        d = ws.cell(row=row + 1, column=2 + i)
        if col.get("year"):
            d.value, d.number_format = datetime.datetime(col["year"], 3, 31), DATE_FMT
        else:
            d.value = col.get("source_name") or ""
        for c in (h, d):
            c.font, c.fill, c.alignment, c.border = (_f(bold=True), _fill(HEAD_BG),
                                                     _a(h="center"), _b())


def _label(ws, row, text, bold=False, fill=None):
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.border = _f(bold=bold), _b()
    if fill:
        c.fill = _fill(fill)
    return c


def _write_input(cell, col: dict, key: str) -> None:
    """One figure read from the accounts, in lakhs - or a flagged 0."""
    v = (col.get("values") or {}).get(key)
    if v is None:
        cell.value = 0
        cell.fill = _fill(UNREAD_BG)
        cell.comment = Comment(UNREAD_NOTE, "ITR Extractor")
        return
    if key == "preference_dividend":
        v = -abs(v)   # the sheet ADDS it: {profit_after_tax}+{preference_dividend}
    cell.value = v / LAKH


def _build_sheet(wb, columns: list) -> None:
    ws = wb.active
    ws.title = "ITR Validation"
    n = max(len(columns), 1)
    ws.column_dimensions["A"].width = 44
    for i in range(n):
        ws.column_dimensions[get_column_letter(2 + i)].width = 15

    title = f"{borrower_name(columns)} (Amt in Lakhs)"
    _period_header(ws, 2, title, columns, merge_title=True)
    ws.freeze_panes = "B4"

    # First pass: every keyed row's position, so formulas can refer forward.
    rowmap, row = {}, 4
    for kind, key, label in ROWS:
        if kind == "blank":
            row += 1
            continue
        if kind == "section":
            _label(ws, row, label, bold=True)
            for i in range(n):
                ws.cell(row=row, column=2 + i).border = _b()
            row += 1
            continue
        rowmap[key] = row
        row += 1

    def refs_for(i):
        L = get_column_letter(2 + i)
        return {k: f"{L}{r}" for k, r in rowmap.items()}

    # Second pass: the P&L and Balance Sheet rows.
    for kind, key, label in ROWS:
        if kind not in ("item", "total"):
            continue
        r = rowmap[key]
        hl = TOTAL_BG if key in C.HIGHLIGHT else None
        _label(ws, r, label or C.LABELS.get(key, key), bold=(kind == "total"), fill=hl)
        for i, col in enumerate(columns):
            cell = ws.cell(row=r, column=2 + i)
            if key in C.TOTAL_FORMULAS:
                cell.value = _formula(C.TOTAL_FORMULAS[key], refs_for(i))
            else:
                _write_input(cell, col, key)
            cell.number_format, cell.alignment, cell.border = AMT_FMT, _a(h="center"), _b()
            cell.font = _f(size=11)
            if hl:
                cell.fill = _fill(hl)

    # RATIOS - live formulas.
    r = row
    _label(ws, r, "RATIOS", bold=True, fill=HEAD_BG)
    for i in range(n):
        ws.cell(row=r, column=2 + i).border = _b()
    r += 1
    ratio_rows = {}
    for name, tmpl, fmt in C.RATIOS:
        _label(ws, r, name, fill=HEAD_BG)
        for i in range(len(columns)):
            c = ws.cell(row=r, column=2 + i, value=_formula(tmpl, refs_for(i)))
            c.number_format, c.alignment, c.border, c.font = fmt, _a(h="center"), _b(), _f(size=11)
        ratio_rows[name] = r
        r += 1

    # DSCR Calculation - linked to the sheet; the analyst keys the EMIs.
    r += 1
    _period_header(ws, r, "DSCR Calculation", columns, merge_title=False)
    ws.cell(row=r + 1, column=1, value="Particular").font = _f(bold=True)
    r += 2
    first = r
    for label, key in (("Net Profit", "profit_after_tax"), ("Depreciation", "depreciation"),
                       ("Interest (Including CC/OD interest)", "interest_finance")):
        _label(ws, r, label)
        for i in range(len(columns)):
            c = ws.cell(row=r, column=2 + i, value=f"={refs_for(i)[key]}")
            c.number_format, c.alignment, c.border = AMT_FMT, _a(h="center"), _b()
        r += 1
    total_a = r
    _label(ws, r, "Total (A)", bold=True)
    for i in range(len(columns)):
        L = get_column_letter(2 + i)
        c = ws.cell(row=r, column=2 + i, value=f"=SUM({L}{first}:{L}{r - 1})")
        c.number_format, c.alignment, c.border, c.font = AMT_FMT, _a(h="center"), _b(), _f(bold=True)
    r += 1
    emi, proposed = r, r + 1
    for label in ("Existing Monthly EMI", "Proposed Loan EMI"):
        _label(ws, r, label)
        for i in range(len(columns)):
            c = ws.cell(row=r, column=2 + i)
            c.number_format, c.alignment, c.border = AMT_FMT, _a(h="center"), _b()
            c.fill = _fill(WARN_BG)   # the analyst's own input
        r += 1
    oblig = r
    _label(ws, r, "Yearly Obligation (B)")
    for i, col in enumerate(columns):
        L = get_column_letter(2 + i)
        m = C.DSCR_PROPOSED_MULTIPLIER[_status(col)]
        c = ws.cell(row=r, column=2 + i, value=f"=SUM({L}{emi}*12+{L}{proposed}*{m})")
        c.number_format, c.alignment, c.border = AMT_FMT, _a(h="center"), _b()
    r += 1
    total_b = r
    _label(ws, r, "Total (B)", bold=True)
    for i in range(len(columns)):
        L = get_column_letter(2 + i)
        c = ws.cell(row=r, column=2 + i, value=f"={L}{oblig}")
        c.number_format, c.alignment, c.border, c.font = AMT_FMT, _a(h="center"), _b(), _f(bold=True)
    r += 1
    _label(ws, r, "DSCR (A/B)", bold=True)
    for i in range(len(columns)):
        L = get_column_letter(2 + i)
        c = ws.cell(row=r, column=2 + i, value=f"=IFERROR({L}{total_a}/{L}{total_b},0)")
        c.number_format, c.alignment, c.border, c.font = "0.00", _a(h="center"), _b(), _f(bold=True)
    r += 1

    # Financial Snap - linked to the rows above.
    r += 2
    _period_header(ws, r, "=A2", columns, merge_title=True)
    r += 2
    for label, tmpl in C.SNAP:
        _label(ws, r, label)
        for i in range(len(columns)):
            c = ws.cell(row=r, column=2 + i, value=_formula(tmpl, refs_for(i)))
            c.number_format, c.alignment, c.border = AMT_FMT, _a(h="center"), _b()
        r += 1
    _label(ws, r, "Ratios", bold=True)
    r += 1
    ratio_fmt = {name: fmt for name, _t, fmt in C.RATIOS}
    for label, source in C.SNAP_RATIO_NAMES:
        _label(ws, r, label)
        for i in range(len(columns)):
            L = get_column_letter(2 + i)
            c = ws.cell(row=r, column=2 + i, value=f"={L}{ratio_rows[source]}")
            c.number_format, c.alignment, c.border = ratio_fmt[source], _a(h="center"), _b()
        r += 1


# ─────────────────────────────────────────────
# SHEET 2 - ANALYSIS
# ─────────────────────────────────────────────

def _build_analysis_sheet(wb, columns: list) -> None:
    """
    Year-on-year trends and red/amber flags (see analysis.py). Separate from
    the ITR Validation sheet so that sheet keeps matching the analyst's own
    template row for row.
    """
    ws = wb.create_sheet("Analysis")
    ws.column_dimensions["A"].width = 44
    n = max(len(columns), 1)
    for i in range(n):
        ws.column_dimensions[get_column_letter(2 + i)].width = 16
    last = get_column_letter(max(1 + n, 4))

    row = 1
    ws.merge_cells(f"A{row}:{last}{row}")
    t = ws.cell(row=row, column=1, value=f"{borrower_name(columns)}  -  Analysis")
    t.font, t.fill, t.alignment = _f(14, True, WHITE), _fill(NAVY), _a(indent=1)
    ws.row_dimensions[row].height = 26
    row += 2

    ws.cell(row=row, column=1, value="Year-on-year growth").font = _f(bold=True, color=NAVY)
    for i, col in enumerate(columns):
        c = ws.cell(row=row, column=2 + i, value=_col_header(col))
        c.font, c.fill, c.alignment = _f(bold=True, color=WHITE), _fill(NAVY), _a(h="center")
    row += 1
    for key, growth in A.trends(columns).items():
        lbl = ws.cell(row=row, column=1, value=A.LABELS[key])
        lbl.alignment, lbl.border = _a(indent=1), _b()
        for i, g in enumerate(growth):
            cell = ws.cell(row=row, column=2 + i)
            if g is None:
                cell.value, cell.font = "-", _f(color="808080")
            else:
                cell.value, cell.number_format = g, PCT_FMT
                cell.font = _f(color=GOOD_GREEN if g >= 0 else BAD_RED)
            cell.alignment, cell.border = _a(h="right"), _b()
        row += 1

    cg = A.cagr(columns)
    ws.cell(row=row, column=1, value="Revenue CAGR (first to last readable year)"
            ).alignment = _a(indent=1)
    c = ws.cell(row=row, column=2, value=cg if cg is not None else "n/a")
    if cg is not None:
        c.number_format = PCT_FMT
    row += 2

    ws.merge_cells(f"A{row}:{last}{row}")
    c = ws.cell(row=row, column=1, value="Flags")
    c.font, c.fill, c.alignment = _f(bold=True, color=NAVY), _fill(SECTION_BG), _a(indent=1)
    row += 1
    fl = A.flags(columns)
    if not fl:
        ws.cell(row=row, column=1, value="No red or amber signals raised."
                ).font = _f(color=GOOD_GREEN)
    for f in fl:
        sev = ws.cell(row=row, column=1,
                      value=f"{f['severity'].upper()}  ·  {f['year'] or '?'}  ·  {f['flag']}")
        sev.font = _f(bold=True, color=BAD_RED if f["severity"] == A.RED else "9C6500")
        ws.merge_cells(f"B{row}:{last}{row}")
        ws.cell(row=row, column=2, value=f["detail"]).font = _f(size=9)
        row += 1
    row += 1
    note = ws.cell(row=row, column=1,
                   value="Thresholds are conventional MSME comfort levels "
                         "(engine/analysis.py THRESHOLDS) - adjust to your credit policy")
    note.font = _f(size=9, italic=True, color="808080")


# ─────────────────────────────────────────────
# SHEET 3 - AUDIT TRAIL
# ─────────────────────────────────────────────

def _build_audit_sheet(wb, columns: list) -> None:
    """
    Where every figure came from, and what could not be verified.

    The point of the arithmetic checks is lost if their results stay inside the
    program: an analyst signing off on these numbers needs to see which
    statement fed which column and which ones proved themselves.
    """
    ws = wb.create_sheet("Audit Trail")
    for col, width in zip("ABCDE", (18, 30, 16, 16, 74)):
        ws.column_dimensions[col].width = width

    row = 1
    t = ws.cell(row=row, column=1, value="Audit Trail")
    t.font = _f(14, True, NAVY)
    row += 2

    for col in columns:
        ws.merge_cells(f"A{row}:E{row}")
        c = ws.cell(row=row, column=1,
                    value=f"{_col_header(col)}  ·  {_status(col)}  ·  {col.get('entity', '')}  ·  "
                          f"{col.get('source_name', '')}"
                          f"{'  ·  scanned' if col.get('scanned') else ''}")
        c.font, c.fill, c.alignment = _f(bold=True, color=WHITE), _fill(NAVY), _a(indent=1)
        row += 1

        # Extraction status - kept off the main sheet so it stays exactly the
        # analyst's format.
        used = col.get("blocks_used") or []
        ok = sum(1 for b in used if b.get("status") == F.VERIFIED)
        excluded = sum(a for _l, a in col.get("unmapped", []))
        for label, value in (
                ("Statements reconciled", f"{ok} of {len(used)}"),
                ("Pages read / in document",
                 f"{col.get('pages_used', 0)} / {col.get('pages_total', 0)}"),
                ("Amount excluded (fits no row)",
                 f"{excluded / LAKH:,.2f} lakhs" if excluded else "none"),
                ("Assumptions made", str(len(col.get("assumptions", []))))):
            ws.cell(row=row, column=1, value=label).font = _f(bold=True)
            v = ws.cell(row=row, column=2, value=value)
            v.font = _f(color=BAD_RED if (label == "Statements reconciled"
                                          and (not used or ok < len(used))) else "000000")
            row += 1

        pages = col.get("page_summary") or {}
        ws.cell(row=row, column=1, value="Pages").font = _f(bold=True)
        ws.cell(row=row, column=2,
                value=", ".join(f"{k}: {v}" for k, v in sorted(pages.items()))).font = _f(size=9)
        row += 1

        vp = col.get("vision_pages") or []
        if vp:
            ws.cell(row=row, column=1, value="Vision calls").font = _f(bold=True)
            ws.merge_cells(f"B{row}:E{row}")
            ws.cell(row=row, column=2,
                    value=f"{len(vp)} page(s) re-read from their images: "
                          + ", ".join(str(p + 1) for p in vp)).font = _f(size=9)
            row += 1

        heads = ("Status", "Statement", "Page", "Total", "Detail")
        for j, head in enumerate(heads, 1):
            hc = ws.cell(row=row, column=j, value=head)
            hc.font, hc.fill, hc.border = _f(bold=True), _fill(SECTION_BG), _b()
        row += 1

        for b in col.get("blocks", []):
            chk = b["check"]
            sc = ws.cell(row=row, column=1, value=b["status"])
            sc.font = _f(bold=True, color={F.VERIFIED: GOOD_GREEN,
                                           F.FAILED: BAD_RED}.get(b["status"], "808080"))
            ws.cell(row=row, column=2,
                    value=f"{b['kind']}  ({b.get('entity') or 'unnamed'})").font = _f(size=9)
            ws.cell(row=row, column=3, value=b["page"] + 1).alignment = _a(h="center")
            tot = ws.cell(row=row, column=4, value=chk.get("printed_total"))
            tot.number_format = "#,##0"
            ws.cell(row=row, column=5,
                    value=chk.get("reason") or "reconciles to the statement's own "
                                               "printed total").font = _f(size=9)
            for cc in range(1, 6):
                ws.cell(row=row, column=cc).border = _b()
            row += 1

        for label, amount in col.get("unmapped", []):
            ws.cell(row=row, column=1, value="EXCLUDED").font = _f(bold=True, color=BAD_RED)
            ws.cell(row=row, column=2, value=label[:60]).font = _f(size=9)
            a = ws.cell(row=row, column=4, value=amount)
            a.number_format = "#,##0"
            a.fill = _fill(WARN_BG)
            ws.cell(row=row, column=5,
                    value="fits no row in this template - NOT included in any "
                          "figure").font = _f(size=9, color=BAD_RED)
            for cc in range(1, 6):
                ws.cell(row=row, column=cc).border = _b()
            row += 1

        for note in col.get("assumptions", []):
            ws.cell(row=row, column=1, value="ASSUMED").font = _f(bold=True, color="9C6500")
            ws.cell(row=row, column=2, value=str(note["label"])[:60]).font = _f(size=9)
            ws.cell(row=row, column=3, value=note["target"]).font = _f(size=9)
            if note["amount"] is not None:
                a = ws.cell(row=row, column=4, value=note["amount"])
                a.number_format = "#,##0"
            ws.cell(row=row, column=5, value=note["note"]).font = _f(size=9)
            for cc in range(1, 6):
                ws.cell(row=row, column=cc).border = _b()
            row += 1

        ignored = col.get("ignored", [])
        if ignored:
            ws.cell(row=row, column=1, value="not counted").font = _f(bold=True, color="808080")
            ws.merge_cells(f"B{row}:E{row}")
            ws.cell(row=row, column=2,
                    value=f"{len(ignored)} subtotal/outcome row(s) "
                          f"(“Total …”, “Profit before tax”, EPS) were read but "
                          f"deliberately excluded - they restate figures already "
                          f"counted, and are re-derived instead"
                    ).font = _f(size=9, color="808080")
            row += 1

        for w in col.get("warnings", []):
            ws.cell(row=row, column=1, value="note").font = _f(bold=True, color="9C6500")
            ws.merge_cells(f"B{row}:E{row}")
            ws.cell(row=row, column=2, value=w).font = _f(size=9)
            row += 1
        row += 1


def get_filename(entity: str, columns: list = None) -> str:
    safe = re.sub(r"[^\w\s-]", "", entity or "ITR").strip().replace(" ", "_") or "ITR"
    span = ""
    if columns:
        years = [c["year"] for c in columns if c.get("year")]
        if years:
            span = f"_{years[0]}" if len(years) == 1 else f"_{years[0]}_{years[-1]}"
    return f"{safe}_ITR_Validation{span}_{datetime.datetime.now():%d%b%Y}.xlsx"
