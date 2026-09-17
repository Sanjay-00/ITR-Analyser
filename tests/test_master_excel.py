"""
The ITR Validation sheet in the combined master format (2026-09-12).

These tests EVALUATE the workbook's own formulas (tests/xlsx_eval.py) - so
they check the Excel the analyst receives, not a Python copy of it - and
compare the results with taxonomy.ratios(), which the app shows.
"""

import datetime
import io
import os
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

from openpyxl import load_workbook                                   # noqa: E402

from engine import generate_excel                                    # noqa: E402
from engine.mapping import taxonomy as T                             # noqa: E402
from engine.mapping import template_config as C                      # noqa: E402
from xlsx_eval import SheetEvaluator                                 # noqa: E402

L = 100000

MASTER_LABELS = [
    "PROFIT AND LOSS ACCOUNT", "Income", "Sales and other Income", "Gross Receipts",
    "Expenses", "Purchases & Raw Material", "Transport operation and admin charges",
    "Employee Costs", "Other Expenses", "Interest and Finance Expenses",
    "Depreciation", "Extra Ordinary Item", "Gross Expenses", "Profit before tax",
    "Provision for tax", "Provision for Deferred Tax", "Profit after tax",
    "Preference dividend (including tax)", "Profit available for equity shareholders",
    "Cash Profit", "BALANCE SHEET", "Liabilities", "Equity Capital", "Preference Shares",
    "Reserves & Surplus", "Secured loan - Asset Financed",
    "Secured loan maturity within one year", "Secured loan - CC/OD",
    "Other Long term liabilities", "Unsecured loans", "Deffered tax liability",
    "Sundry Creditors", "Current Liabilities and Provision", "Total of Liabilities",
    "Assets", "Fixed Assets", "Intangible Assets", "Investments", "Deffered Tax Assets",
    "Debtors / Receivables", "Current assets, Loans and Advances",
    "Long Term Loan & Advance", "Other Non Current Assets", "Total of Assets",
    "RATIOS", "Return on Capital Employed", "Return On Share Holders Fund",
    "PAT / Income (%)", "PAT / Assets employed (%)",
    "Long Term Debt / Equity (inclusive q/e)", "Total Debt / Equity(inclusive q/e)",
    "Interest Coverage", "Current Ratio", "DSCR", "Debtor Days", "Creditor Days",
    "Turnover and Other Income", "PAT", "Cash Profit", "Networth",
    "Secured Loans - Long Term", "Secured Loans – CC / OD",
    "Current Liabilities and Provisions", "Fixed Assets",
    "Current Assets , ST Loans & Advances ", "Ratios", "PAT/Income  (%)",
    "Long Term Debt / Equity (including Q/E)", "TOL/TNW (including Q/E)",
    "Interest Coverage", "Current Ratio", "DSCR", "Debtor Days", "Creditor Days",
    "Year-on-year growth", "Sales and other Income", "Gross Expenses",
    "Profit before tax", "Profit after tax", "Cash Profit", "Total of Assets",
    "Total of Liabilities",
]


def _column(year, buckets, have_pl=True, have_bs=True, provisional=False):
    v = T.compute(buckets, have_pl=have_pl, have_bs=have_bs)
    return {"year": year, "entity": "Test Traders", "values": v, "ratios": T.ratios(v),
            "provisional": provisional, "blocks": [],
            "blocks_used": [{"status": "verified", "kind": "profit_loss"}],
            "warnings": [], "unmapped": [], "assumptions": [], "ignored": [],
            "page_summary": {}, "vision_pages": [], "source_name": f"ITR {year}.pdf",
            "pages_used": 5, "pages_total": 20, "scanned": False, "book_profit": None}


BUCKETS = {
    "sales_other_income": 500 * L, "purchases": 200 * L, "transport_admin": 120 * L,
    "employee_costs": 40 * L, "other_expenses": 10 * L, "interest_finance": 30 * L,
    "depreciation": 50 * L, "provision_tax": 12 * L, "provision_deferred_tax": 3 * L,
    "equity_capital": 200 * L, "reserves": 50 * L, "secured_loan_asset_financed": 300 * L,
    "secured_loan_ccod": 40 * L, "unsecured_loans": 20 * L, "sundry_creditors": 60 * L,
    "current_liabilities": 25 * L, "fixed_assets": 450 * L, "debtors": 90 * L,
    "current_assets": 155 * L,
}


@pytest.fixture(scope="module")
def sheet():
    cols = [_column(2024, BUCKETS),
            _column(2025, {}, have_pl=False, have_bs=False, provisional=True)]
    ws = load_workbook(io.BytesIO(generate_excel(cols)))["ITR Validation"]
    return ws, SheetEvaluator(ws), cols


def _row(ws, label, start=1):
    for r in range(start, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == label:
            return r
    raise AssertionError(f"no row {label!r}")


def test_layout_is_exactly_the_master_format(sheet):
    ws, ev, _ = sheet
    labels = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
    labels = [l for l in labels if l and l != "=A2"]
    assert labels == MASTER_LABELS
    assert ws["A2"].value == "Test Traders (Amt in Lakhs)"
    assert ws["B2"].value == "Audited" and ws["C2"].value == "Provisional"
    assert ws["B3"].value == datetime.datetime(2024, 3, 31)


def test_profit_and_totals_are_live_formulas_with_tax_subtracted(sheet):
    ws, ev, cols = sheet
    pbt, pat = _row(ws, "Profit before tax"), _row(ws, "Profit after tax")
    assert str(ws.cell(row=pat, column=2).value).startswith("=")
    assert ev.value(f"B{pbt}") == pytest.approx(500 - 450)
    assert ev.value(f"B{pat}") == pytest.approx(50 - 12 - 3)
    assert ev.value(f"B{_row(ws, 'Cash Profit')}") == pytest.approx(35 + 50)
    assert ev.value(f"B{_row(ws, 'Total of Liabilities')}") == pytest.approx(695)
    assert ev.value(f"B{_row(ws, 'Total of Assets')}") == pytest.approx(695)


def test_excel_ratios_equal_the_python_ratios(sheet):
    ws, ev, cols = sheet
    py = cols[0]["ratios"]
    for name in C.RATIO_NAMES:
        got = ev.value(f"B{_row(ws, name)}")
        want = py[name] if py[name] is not None else 0
        assert got == pytest.approx(want), name
    # Spot-check against hand arithmetic from the analysts' definitions.
    assert py["Current Ratio"] == pytest.approx((90 + 155) / (60 + 25 + 40))
    assert py["Debtor Days"] == pytest.approx(90 / 500 * 365)
    assert py["Long Term Debt / Equity (inclusive q/e)"] == pytest.approx(300 / (250 + 20))


def test_unread_year_is_zero_but_flagged_and_formulas_still_work(sheet):
    ws, ev, _ = sheet
    r = _row(ws, "Sales and other Income")
    cell = ws.cell(row=r, column=3)
    assert cell.value == 0
    assert cell.fill.fgColor.rgb.endswith("FFC7CE")
    assert cell.comment is not None and "Not read" in cell.comment.text
    # Every formula must still evaluate on an all-zero column - each ratio
    # divides by zero there and must fall back to 0 through its IFERROR.
    for label in ["Profit after tax", "Total of Assets"] + C.RATIO_NAMES:
        assert ev.value(f"C{_row(ws, label)}") == pytest.approx(0), label


def test_snap_repeats_the_sheet(sheet):
    ws, ev, _ = sheet
    snap_start = _row(ws, "Turnover and Other Income") - 1
    assert ev.value(f"B{_row(ws, 'Networth', snap_start)}") == pytest.approx(250)
    assert ev.value(f"B{_row(ws, 'Current Assets , ST Loans & Advances ', snap_start)}") \
        == pytest.approx(245)
    assert ev.value(f"B{_row(ws, 'Debtor Days', _row(ws, 'Ratios', snap_start))}") \
        == pytest.approx(90 / 500 * 365)


def test_growth_block_closes_the_sheet(sheet):
    """The sheet ends with year-on-year growth, as LIVE formulas off the rows
    above (the analyst's own "ITR - Borrower D" workbook, 2026-09-16):
    correct a figure and its growth follows. The first year has nothing to
    compare against and shows "-"."""
    ws, ev, _ = sheet
    head = _row(ws, "Year-on-year growth")
    sales = _row(ws, "Sales and other Income", head)
    assert ws.cell(row=sales, column=2).value == "-"
    formula = ws.cell(row=sales, column=3).value
    assert formula.startswith("=(C") and "ABS(B" in formula
    # FY2025 in the fixture is unread, so growth off a 500 base is -100%.
    assert ev.value(f"C{sales}") == pytest.approx(-1.0)
    assert ws.cell(row=sales, column=3).number_format.startswith("0.0%")


def test_each_figure_shows_the_lines_it_adds_up():
    """
    Transparency: a figure made of several statement lines is written as
    their live sum ("=1.2+0.1+0.05"), with a note naming each line, its page
    and file - and the Audit Trail lists them all. A single-line figure keeps
    its plain value but still carries the note.
    """
    col = _column(2024, BUCKETS)
    col["sources"] = {
        "secured_loan_ccod": [
            {"label": "SBI Cash Credit", "amount": 25 * L, "page": 6, "file": "ITR 2024.pdf"},
            {"label": "HDFC OD", "amount": 15 * L, "page": 6, "file": "ITR 2024.pdf"}],
        "purchases": [
            {"label": "Purchases", "amount": 200 * L, "page": 7, "file": "ITR 2024.pdf"}],
    }
    wb = load_workbook(io.BytesIO(generate_excel([col])))
    ws = wb["ITR Validation"]
    ev = SheetEvaluator(ws)
    cc = ws.cell(row=_row(ws, "Secured loan - CC/OD"), column=2)
    assert cc.value == "=25+15"
    assert ev.value(cc.coordinate) == pytest.approx(40)
    assert "SBI Cash Credit" in cc.comment.text and "p6" in cc.comment.text
    pur = ws.cell(row=_row(ws, "Purchases & Raw Material"), column=2)
    assert pur.value == pytest.approx(200) and "Purchases" in pur.comment.text
    # Totals still work on top of the summed cells.
    assert ev.value(f"B{_row(ws, 'Total of Liabilities')}") == pytest.approx(695)
    audit = [c.value for r in wb["Audit Trail"].iter_rows() for c in r]
    assert "HDFC OD" in audit and "Secured loan - CC/OD" in audit


def test_breakdown_not_written_when_it_does_not_add_up():
    """Never show a sum that disagrees with the figure itself."""
    col = _column(2024, BUCKETS)
    col["sources"] = {"secured_loan_ccod": [
        {"label": "SBI Cash Credit", "amount": 25 * L, "page": 6, "file": "x.pdf"}]}
    ws = load_workbook(io.BytesIO(generate_excel([col])))["ITR Validation"]
    cc = ws.cell(row=_row(ws, "Secured loan - CC/OD"), column=2)
    assert cc.value == pytest.approx(40) and cc.comment is None
