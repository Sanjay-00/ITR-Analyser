"""
Regression suite  -  runs against an external folder of real sample ITR
bundles, exercising the actual shipped pipeline (parser.spread ->
excel_generator.generate_excel).

    ITR_TEST_DIR="C:\\path\\to\\sample\\borrowers" pytest tests/

Without ITR_TEST_DIR set, every test skips cleanly. Real ITRs carry PAN,
Aadhaar, addresses and bank account numbers - they must NEVER be committed to
this repo, which is exactly why the samples live in an env-var'd folder
outside it rather than in a fixtures directory.

Expected layout under ITR_TEST_DIR: one subdirectory per borrower, each
holding that borrower's ITR PDFs (one per financial year) - the same shape
spread() takes in production.

These are internal-consistency checks, not golden-value comparisons: they
assert that whatever was extracted agrees with the pipeline's own arithmetic
and its own invariants (None vs 0, mapping conservation, sheet shape). That
catches the failure mode that actually matters - a statement misread, or a
value silently defaulted - without needing a hand-keyed expected figure per
sample. Run offline (no API key): this checks the rule-based/OCR path, which
is what every column falls back to when Vision is off or unavailable.
"""

import glob
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from engine.extract import financials as F      # noqa: E402
from engine import generate_excel, spread       # noqa: E402

TEST_DIR = os.getenv("ITR_TEST_DIR", "")


def _borrower_dirs():
    if not TEST_DIR or not os.path.isdir(TEST_DIR):
        return []
    return sorted(d for d in glob.glob(os.path.join(TEST_DIR, "*"))
                  if os.path.isdir(d) and glob.glob(os.path.join(d, "*.pdf")))


BORROWERS = _borrower_dirs()

pytestmark = pytest.mark.skipif(
    not BORROWERS,
    reason="Set ITR_TEST_DIR to a folder of per-borrower ITR PDF subfolders "
           "to run these tests",
)


@pytest.fixture(scope="module")
def spread_result():
    """Spread every borrower once (offline, no API key); tests read from this."""
    return {d: spread(sorted(glob.glob(os.path.join(d, "*.pdf"))))
            for d in BORROWERS}


# ── Per-column invariants ───────────────────────────────────────────

@pytest.mark.parametrize("path", BORROWERS, ids=os.path.basename)
def test_every_verified_statement_actually_balances(spread_result, path):
    """
    A block marked VERIFIED must genuinely have reconciled - nothing
    downstream re-checks this, the status is trusted as-is.
    """
    for col in spread_result[path]:
        for b in col.get("blocks", []):
            if b["status"] == F.VERIFIED:
                assert b["check"]["balanced"], (
                    f"{col.get('source_name')}: block on page {b['page'] + 1} "
                    f"marked VERIFIED but its own check reports not balanced")


@pytest.mark.parametrize("path", BORROWERS, ids=os.path.basename)
def test_mapping_conserves_every_rupee(spread_result, path):
    """Nothing a statement printed should vanish or double up during mapping."""
    for col in spread_result[path]:
        mc = col.get("mapping_check")
        if mc is not None:
            assert mc["ok"], f"{col.get('source_name')}: mapping lost {mc['diff']:,}"


@pytest.mark.parametrize("path", BORROWERS, ids=os.path.basename)
def test_unread_statement_is_none_not_zero(spread_result, path):
    """
    A column with no usable P&L (or no usable balance sheet) must report
    None for that statement's rows - "Check ITR" - never a column of zeros,
    which reads as "this business earned/owns nothing" rather than "we
    could not read this".
    """
    for col in spread_result[path]:
        used = col.get("blocks_used") or []
        have_pl = any(b["kind"] == F.PROFIT_LOSS for b in used)
        have_bs = any(b["kind"] == F.BALANCE_SHEET for b in used)
        vals = col.get("values") or {}
        if not have_pl:
            assert vals.get("profit_before_tax") is None, (
                f"{col.get('source_name')}: no P&L read, but "
                f"profit_before_tax is not None")
        if not have_bs:
            assert vals.get("total_assets") is None, (
                f"{col.get('source_name')}: no balance sheet read, but "
                f"total_assets is not None")


@pytest.mark.parametrize("path", BORROWERS, ids=os.path.basename)
def test_entity_present_when_statements_were_used(spread_result, path):
    """Every column that actually spread a statement must name who it is."""
    for col in spread_result[path]:
        if col.get("blocks_used"):
            assert col.get("entity"), (
                f"{col.get('source_name')}: statements were used but no "
                f"entity name was captured")


# ── Excel output ─────────────────────────────────────────────────────

@pytest.mark.parametrize("path", BORROWERS, ids=os.path.basename)
def test_excel_generates(spread_result, path):
    """The writer survives whatever this borrower's columns produced,
    including None ('Check ITR') fields and a NIL-return year."""
    from openpyxl import load_workbook
    import io

    columns = spread_result[path]
    data = generate_excel(columns)
    assert data, "generate_excel returned no bytes"

    wb = load_workbook(io.BytesIO(data))
    assert "ITR Validation" in wb.sheetnames
    assert "Audit Trail" in wb.sheetnames
