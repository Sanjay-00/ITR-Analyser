"""
xlsx_eval.py  -  evaluate the formulas this project writes into its own
workbook, for tests.

The ITR Validation sheet holds live Excel formulas (totals, profit, ratios),
and openpyxl never computes them - a freshly written file has no cached
values. Rather than re-deriving each figure in Python (which would test a
copy of the formula, not the formula), this evaluates the sheet's actual
formula text.

Deliberately small: it understands exactly the grammar excel_generator
emits - numbers, cell references, A1:B9 ranges, + - * /, parentheses,
SUM(...) and IFERROR(x, y). Anything else raises, so a new formula shape
fails loudly in the tests instead of being silently mis-evaluated.
"""

import re

from openpyxl.utils import column_index_from_string, get_column_letter

_TOKEN = re.compile(
    r"\s*(?:(?P<num>\d+(?:\.\d+)?)|(?P<ref>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)"
    r"|(?P<func>SUM|IFERROR)\s*\(|(?P<op>[+\-*/(),])|(?P<str>\"[^\"]*\"))")


class _Error(Exception):
    """An Excel-style error value (#DIV/0!, #VALUE!)."""


def _tokens(text):
    pos, out = 0, []
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if not m or m.end() == pos:
            if text[pos:].strip() == "":
                break
            raise ValueError(f"unsupported formula text at {text[pos:]!r}")
        kind = m.lastgroup
        out.append((kind, m.group(kind)))
        pos = m.end()
    return out


class SheetEvaluator:
    def __init__(self, ws):
        self.ws = ws
        self._memo = {}

    def value(self, coord):
        """Computed value of one cell: a number, a string, or None."""
        coord = coord.replace("$", "")
        if coord in self._memo:
            return self._memo[coord]
        raw = self.ws[coord].value
        if isinstance(raw, str) and raw.startswith("="):
            try:
                val = self._eval(raw[1:])
            except _Error:
                val = "#ERROR"
        else:
            val = raw
        self._memo[coord] = val
        return val

    # ── recursive descent over the token list ─────────────────────
    def _eval(self, text):
        # Re-entrant: a formula's cell references evaluate OTHER formulas in
        # the middle of this parse, so the parser state is saved and restored
        # around each one rather than overwritten.
        saved = (getattr(self, "_toks", None), getattr(self, "_i", 0))
        self._toks, self._i = _tokens(text), 0
        try:
            v = self._expr()
            if self._i != len(self._toks):
                raise ValueError(f"trailing tokens in {text!r}")
            return v
        finally:
            self._toks, self._i = saved

    def _peek(self):
        return self._toks[self._i] if self._i < len(self._toks) else (None, None)

    def _take(self):
        t = self._toks[self._i]
        self._i += 1
        return t

    def _expr(self):
        v = self._term()
        while self._peek()[1] in ("+", "-"):
            op = self._take()[1]
            r = self._term()
            v = _num(v) + _num(r) if op == "+" else _num(v) - _num(r)
        return v

    def _term(self):
        v = self._factor()
        while self._peek()[1] in ("*", "/"):
            op = self._take()[1]
            r = self._factor()
            if op == "*":
                v = _num(v) * _num(r)
            else:
                d = _num(r)
                if d == 0:
                    raise _Error("#DIV/0!")
                v = _num(v) / d
        return v

    def _factor(self):
        kind, tok = self._take()
        if kind == "num":
            return float(tok)
        if kind == "str":
            return tok[1:-1]
        if kind == "op" and tok in ("-", "+"):
            f = self._factor()
            return -_num(f) if tok == "-" else _num(f)
        if kind == "op" and tok == "(":
            v = self._expr()
            self._expect(")")
            return v
        if kind == "ref":
            if ":" in tok:
                raise ValueError("a range is only valid inside SUM")
            return self._cell(tok)
        if kind == "func" and tok.startswith("SUM"):
            total = 0.0
            while True:
                k, t = self._peek()
                if k == "ref" and ":" in t:
                    self._take()
                    total += sum(_num_or_zero(self._cell(c)) for c in _expand(t))
                else:
                    total += _num_or_zero(self._expr())
                if self._peek()[1] == ",":
                    self._take()
                    continue
                break
            self._expect(")")
            return total
        if kind == "func" and tok.startswith("IFERROR"):
            start = self._i
            try:
                v = self._expr()
                self._expect(",")
                # skip the fallback expression
                depth = 0
                while True:
                    k, t = self._peek()
                    if _opens(k, t):
                        depth += 1
                    elif t == ")":
                        if depth == 0:
                            break
                        depth -= 1
                    self._take()
                self._expect(")")
                return v
            except _Error:
                # re-scan to the fallback after the top-level comma
                self._i = start
                depth = 0
                while True:
                    k, t = self._take()
                    if _opens(k, t):
                        depth += 1
                    elif t == ")":
                        depth -= 1
                    elif t == "," and depth == 0:
                        break
                fb = self._expr()
                self._expect(")")
                return fb
        raise ValueError(f"unexpected token {tok!r}")

    def _expect(self, sym):
        k, t = self._take()
        if t != sym:
            raise ValueError(f"expected {sym!r}, got {t!r}")

    def _cell(self, coord):
        v = self.value(coord)
        if v == "#ERROR":
            raise _Error("#REF")
        return v


def _opens(kind, tok):
    """Does this token open a bracket? A function token ("SUM(", "IFERROR(")
    carries its own "(" - missing that sent the bracket count negative and
    ran off the end whenever an IFERROR caught an error around a SUM (found
    on a real unread column: TOL/TNW dividing by an all-zero equity)."""
    return tok == "(" or kind == "func"


def _num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    raise _Error("#VALUE!")


def _num_or_zero(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def _expand(rng):
    a, b = rng.replace("$", "").split(":")
    ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups()
    cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
    cols = range(column_index_from_string(ca), column_index_from_string(cb) + 1)
    return [f"{get_column_letter(c)}{r}" for c in cols
            for r in range(int(ra), int(rb) + 1)]
